import pytest
from docx import Document
from docx.oxml.ns import qn

from hanzi_trainer import (
    DATA_FILE,
    MODE_FONT,
    MODE_PINYIN,
    PINYIN_FONT,
    PLAIN_FONT,
    RUBY_FONT,
    apply_exceptions,
    build_answer_key,
    build_document,
    character_levels,
    decode_setup,
    document_bytes,
    encode_setup,
    extract_text,
    load_known_chars,
    load_levels,
    main,
    parse_selections,
    passage_report,
    pinyin_per_character,
    projector_page,
    report_csv,
    worksheet_bundle,
)


def test_levels_available():
    levels = load_levels(DATA_FILE)
    assert levels == {
        "YCT": 4,
        "HSK2.0": 6,
        "HSK3.0": 7,
        "HSK3.0-2026": 7,
    }


@pytest.mark.parametrize(
    "selections, expected",
    [
        ({"YCT": 1}, 99),
        ({"YCT": 3}, 353),
        ({"HSK3.0": 2}, 598),
    ],
)
def test_up_to_level_counts(selections, expected):
    # Regression values for the bundled data/levels.csv.
    assert len(load_known_chars(DATA_FILE, selections)) == expected


def test_levels_are_cumulative():
    lower = load_known_chars(DATA_FILE, {"YCT": 2})
    higher = load_known_chars(DATA_FILE, {"YCT": 3})
    assert lower < higher


def test_multiple_systems_union():
    a = load_known_chars(DATA_FILE, {"YCT": 2})
    b = load_known_chars(DATA_FILE, {"HSK3.0": 1})
    both = load_known_chars(DATA_FILE, {"YCT": 2, "HSK3.0": 1})
    assert both == a | b


def test_unknown_system_raises():
    with pytest.raises(ValueError, match="Unknown word list"):
        load_known_chars(DATA_FILE, {"HSK": 3})


def test_parse_selections():
    assert parse_selections("YCT=3, HSK3.0=1") == {"YCT": 3, "HSK3.0": 1}
    with pytest.raises(ValueError):
        parse_selections("YCT")


def test_document_fonts_follow_known_characters():
    # 你 is in YCT 1; 龙 is not.
    known = load_known_chars(DATA_FILE, {"YCT": 1})
    assert "你" in known and "龙" not in known
    doc = build_document("你龙\n你", known, mode=MODE_FONT)
    paragraphs = doc.paragraphs
    assert len(paragraphs) == 2
    first = paragraphs[0].runs
    assert [r.text for r in first] == ["你", "龙"]
    assert first[0].font.name == PLAIN_FONT
    assert first[1].font.name == PINYIN_FONT
    for run in first:
        east_asia = run._element.rPr.rFonts.get(qn("w:eastAsia"))
        assert east_asia == run.font.name


def test_consecutive_same_font_characters_share_a_run():
    known = load_known_chars(DATA_FILE, {"YCT": 1})
    doc = build_document("你好我", known, mode=MODE_FONT)
    assert len(doc.paragraphs[0].runs) == 1


def test_document_bytes_roundtrip(tmp_path):
    out = tmp_path / "w.docx"
    out.write_bytes(document_bytes(build_document("你好", set(), mode=MODE_FONT)))
    assert Document(str(out)).paragraphs[0].text == "你好"


# --- pinyin (ruby) mode -------------------------------------------------------


def _ruby_pairs(doc):
    """[(base, pinyin)] for every ruby annotation in document order."""
    pairs = []
    for ruby in doc.element.body.iter(qn("w:ruby")):
        rt = "".join(t.text for t in ruby.find(qn("w:rt")).iter(qn("w:t")))
        base = "".join(t.text for t in ruby.find(qn("w:rubyBase")).iter(qn("w:t")))
        pairs.append((base, rt))
    return pairs


def test_pinyin_per_character_is_aligned():
    assert pinyin_per_character("ab你c") == ["", "", "nǐ", ""]
    assert pinyin_per_character("") == []


def test_pinyin_uses_phrase_context():
    assert pinyin_per_character("银行") == ["yín", "háng"]
    assert pinyin_per_character("他行不行") == ["tā", "xíng", "bù", "xíng"]
    assert pinyin_per_character("长大了") == ["zhǎng", "dà", "le"]


def test_pinyin_mode_annotates_only_unknown_hanzi():
    known = load_known_chars(DATA_FILE, {"YCT": 1})
    doc = build_document("你龙", known)  # default mode is pinyin
    assert _ruby_pairs(doc) == [("龙", "lóng")]
    # The known character is an ordinary plain run.
    plain = doc.paragraphs[0].runs[0]
    assert plain.text == "你" and plain.font.name == PLAIN_FONT


def test_pinyin_mode_skips_known_punctuation_and_ascii():
    known = load_known_chars(DATA_FILE, {"YCT": 1})
    doc = build_document("你好,Hello 123。", known)
    assert _ruby_pairs(doc) == []
    assert doc.paragraphs[0].text == "你好,Hello 123。"


def test_pinyin_mode_one_paragraph_per_line():
    doc = build_document("龙\n\n龙", set())
    assert len(doc.paragraphs) == 3
    assert _ruby_pairs(doc) == [("龙", "lóng"), ("龙", "lóng")]


def test_ruby_properties_and_fonts():
    doc = build_document("龙", set(), size_pt=30, ruby_pt=10)
    ruby = next(doc.element.body.iter(qn("w:ruby")))
    props = ruby.find(qn("w:rubyPr"))
    assert props.find(qn("w:hps")).get(qn("w:val")) == "20"
    assert props.find(qn("w:hpsBaseText")).get(qn("w:val")) == "60"
    rt_fonts = ruby.find(qn("w:rt")).find(".//" + qn("w:rFonts"))
    base_fonts = ruby.find(qn("w:rubyBase")).find(".//" + qn("w:rFonts"))
    assert rt_fonts.get(qn("w:eastAsia")) == RUBY_FONT
    assert base_fonts.get(qn("w:eastAsia")) == PLAIN_FONT


def test_font_names_with_special_characters_are_escaped(tmp_path):
    doc = build_document("龙", set(), plain_font='A"B&<C', ruby_font="X'Y>")
    out = tmp_path / "w.docx"
    out.write_bytes(document_bytes(doc))
    assert _ruby_pairs(Document(str(out))) == [("龙", "lóng")]


def test_unknown_mode_raises():
    with pytest.raises(ValueError, match="Unknown mode"):
        build_document("你", set(), mode="nope")


def test_cli_writes_a_worksheet_in_each_mode(tmp_path):
    passage = tmp_path / "p.txt"
    passage.write_text("你好,我是龙。\n", encoding="utf-8")
    for mode in (MODE_PINYIN, MODE_FONT):
        out = tmp_path / f"{mode}.docx"
        assert main(["YCT=1", str(passage), "-o", str(out), "--mode", mode]) == 0
        assert Document(str(out)).paragraphs
    assert _ruby_pairs(Document(str(tmp_path / "pinyin.docx"))) == [("龙", "lóng")]
    assert _ruby_pairs(Document(str(tmp_path / "font.docx"))) == []


# --- exceptions and answer key -------------------------------------------------


def test_exceptions_work_in_both_directions():
    known = load_known_chars(DATA_FILE, {"YCT": 1})
    assert "你" in known and "龙" not in known
    adjusted = apply_exceptions(known, show_pinyin="你", hide_pinyin="龙")
    assert "你" not in adjusted and "龙" in adjusted
    assert known == load_known_chars(DATA_FILE, {"YCT": 1})  # input untouched


def test_exceptions_ignore_non_chinese_and_empty():
    known = {"你"}
    assert apply_exceptions(known, "abc 1,", "") == known


def test_exception_in_both_lists_raises():
    with pytest.raises(ValueError, match="both exception lists: 龙"):
        apply_exceptions(set(), show_pinyin="龙猫", hide_pinyin="龙")


def test_exceptions_change_the_worksheet():
    known = apply_exceptions(
        load_known_chars(DATA_FILE, {"YCT": 1}), show_pinyin="你", hide_pinyin="龙"
    )
    pairs = _ruby_pairs(build_document("你龙", known))
    assert pairs == [("你", "nǐ")]


def test_answer_key_has_pinyin_on_every_character():
    key = build_answer_key("你龙,好。")
    assert _ruby_pairs(key) == [("你", "nǐ"), ("龙", "lóng"), ("好", "hǎo")]


def test_bundle_holds_student_and_teacher_copies(tmp_path):
    import io
    import zipfile

    known = load_known_chars(DATA_FILE, {"YCT": 1})
    z = zipfile.ZipFile(io.BytesIO(worksheet_bundle("你龙", known)))
    assert sorted(z.namelist()) == ["student_worksheet.docx", "teacher_answer_key.docx"]
    z.extractall(tmp_path)
    student = Document(str(tmp_path / "student_worksheet.docx"))
    teacher = Document(str(tmp_path / "teacher_answer_key.docx"))
    assert _ruby_pairs(student) == [("龙", "lóng")]
    assert _ruby_pairs(teacher) == [("你", "nǐ"), ("龙", "lóng")]


def test_cli_exceptions_and_answer_key(tmp_path):
    passage = tmp_path / "p.txt"
    passage.write_text("你龙\n", encoding="utf-8")
    out, key = tmp_path / "s.docx", tmp_path / "k.docx"
    assert main(["YCT=1", str(passage), "-o", str(out), "--show-pinyin", "你",
                 "--hide-pinyin", "龙", "--answer-key", str(key)]) == 0
    assert _ruby_pairs(Document(str(out))) == [("你", "nǐ")]
    assert len(_ruby_pairs(Document(str(key)))) == 2


# --- file upload ----------------------------------------------------------------


def test_extract_text_reads_utf8_bom_and_gb18030():
    assert extract_text("a.txt", "你好".encode("utf-8")) == "你好"
    assert extract_text("a.TXT", "你好".encode("utf-8-sig")) == "你好"
    assert extract_text("a.txt", "你好,我是老师".encode("gb18030")) == "你好,我是老师"


def test_extract_text_from_plain_docx():
    doc = Document()
    doc.add_paragraph("你好")
    doc.add_paragraph("我是老师")
    assert extract_text("a.docx", document_bytes(doc)) == "你好\n我是老师"


def test_extract_text_reads_tables():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格里的字"
    assert "表格里的字" in extract_text("a.docx", document_bytes(doc))


def test_extract_text_roundtrips_our_own_worksheet():
    # The pinyin (ruby text) must not come back as extra letters.
    sheet = build_document("龙你好\n明天", {"你"})
    assert _ruby_pairs(sheet)  # it does contain pinyin
    assert extract_text("w.docx", document_bytes(sheet)) == "龙你好\n明天"


def test_extract_text_rejects_bad_files():
    with pytest.raises(ValueError, match="txt or .docx"):
        extract_text("a.pdf", b"%PDF")
    with pytest.raises(ValueError, match="txt or .docx"):
        extract_text("noextension", b"x")
    with pytest.raises(ValueError, match="Word file"):
        extract_text("a.docx", b"not a zip")
    with pytest.raises(ValueError, match="encoding"):
        extract_text("a.txt", b"\xff\xfe\xfa\x81")


# --- projector page ---------------------------------------------------------------


def test_projector_page_marks_known_pinyin_as_hidden():
    page = projector_page("你龙", {"你"})
    assert '<ruby>你<rt class="k">nǐ</rt></ruby>' in page  # known: hidden until revealed
    assert "<ruby>龙<rt>lóng</rt></ruby>" in page  # unknown: always shown


def test_projector_page_has_fullscreen_and_reveal_controls():
    page = projector_page("你", set())
    for needle in ('id="fs"', "requestFullscreen", 'id="rv"', "Show all pinyin"):
        assert needle in page


def test_projector_page_escapes_and_keeps_lines():
    page = projector_page("你<b>\n\n龙", set(), plain_font='A";}x', ruby_font="Arial")
    assert "<b>" not in page.split('id="text"')[1].split("</div>")[0]
    assert "&lt;b&gt;" in page
    assert page.count("<br>") == 2
    assert '"Ax"' in page  # unsafe characters stripped from the font name


# --- class set-ups in the URL ---------------------------------------------------------

LEVELS = load_levels(DATA_FILE)


def test_encode_leaves_out_defaults_and_the_passage():
    assert encode_setup({"YCT": 3}) == {"lists": "YCT:3"}
    assert encode_setup({}) == {"lists": "-"}
    params = encode_setup({"YCT": 2, "HSK3.0": 1}, show="老师a", hide="北京", mode="font", size_pt=36,
                          plain_font="STKaiti")
    assert params == {"lists": "YCT:2,HSK3.0:1", "show": "师老", "hide": "京北", "mode": "font",
                      "size": "36", "cfont": "STKaiti"}


def test_setup_roundtrips():
    original = dict(selections={"YCT": 4, "HSK2.0": 2}, show="老", hide="龙京", mode="font",
                    size_pt=40, ruby_pt=9, plain_font="STKaiti", ruby_font="Calibri")
    params = encode_setup(original["selections"], original["show"], original["hide"], original["mode"],
                          original["size_pt"], original["ruby_pt"], original["plain_font"], original["ruby_font"])
    decoded = decode_setup(params, LEVELS)
    assert decoded["selections"] == original["selections"]
    assert decoded["show"] == "老" and decoded["hide"] == "京龙"
    assert (decoded["mode"], decoded["size_pt"], decoded["ruby_pt"]) == ("font", 40, 9)
    assert (decoded["plain_font"], decoded["ruby_font"]) == ("STKaiti", "Calibri")


def test_decode_ignores_garbage_and_clamps():
    decoded = decode_setup(
        {"lists": "YCT:99,NOPE:1,HSK3.0:x,HSK2.0:0", "mode": "evil", "size": "9999", "psize": "abc",
         "show": "abc<script>老", "cfont": "  "},
        LEVELS,
    )
    assert decoded == {"selections": {"YCT": LEVELS["YCT"], "HSK2.0": 1}, "show": "老", "size_pt": 48}


def test_decode_empty_and_none_selection():
    assert decode_setup({}, LEVELS) == {}
    assert decode_setup({"lists": "-"}, LEVELS) == {"selections": {}}
    assert decode_setup({"lists": "nonsense"}, LEVELS) == {}


# --- passage report ---------------------------------------------------------------------


def test_report_counts_and_percentages():
    report = passage_report("你你龙,abc。", {"你"})
    assert (report["total"], report["known_total"]) == (3, 2)
    assert (report["distinct"], report["known_distinct"]) == (2, 1)
    assert (report["pct_text"], report["pct_distinct"]) == (67, 50)
    assert [(r["char"], r["count"], r["known"]) for r in report["rows"]] == [
        ("你", 2, True),
        ("龙", 1, False),
    ]


def test_report_lists_every_reading_of_a_character():
    report = passage_report("银行\n他行不行", set())
    row = next(r for r in report["rows"] if r["char"] == "行")
    assert row["pinyin"] == "háng / xíng" and row["count"] == 3


def test_report_for_text_without_chinese():
    report = passage_report("hello 123", set())
    assert report["rows"] == [] and report["pct_text"] is None and report["pct_distinct"] is None


def test_character_levels_use_the_lowest_level_and_first_list():
    levels = character_levels(DATA_FILE, ["YCT", "HSK3.0"])
    assert levels["你"] == "YCT 1"  # in both lists at level 1: the first list given wins
    assert character_levels(DATA_FILE, []) == {}
    assert all(label.split()[0] in ("YCT", "HSK3.0") for label in levels.values())


def test_report_levels_flow_into_rows_and_csv():
    levels = character_levels(DATA_FILE, ["YCT"])
    report = passage_report("你龙", {"你"}, levels)
    assert report["rows"][0]["level"] == "YCT 1"
    data = report_csv(report)
    assert data.startswith(b"\xef\xbb\xbf")  # BOM so Excel reads the Chinese correctly
    rows = list(__import__("csv").reader(data.decode("utf-8-sig").splitlines()))
    assert rows[0] == ["Character", "Pinyin", "Times in passage", "Status", "First taught at"]
    assert rows[1][:5] == ["你", "nǐ", "1", "Known", "YCT 1"]
    assert rows[2][3] == "To learn"


# --- worksheet layout options -------------------------------------------------------------


def _lines(doc):
    """Paragraph texts including characters inside pinyin annotations."""
    return extract_text("x.docx", document_bytes(doc)).split("\n")


def test_title_and_name_line_come_first():
    doc = build_document("龙", set(), title="第一课", name_line=True)
    texts = _lines(doc)
    assert texts[0] == "第一课" and texts[1].startswith("Name:") and "Date:" in texts[1]
    assert texts[2] == "龙"
    heading = doc.paragraphs[0]
    assert heading.alignment == 1  # centred
    assert heading.runs[0].bold is True


def test_no_title_or_name_line_by_default():
    doc = build_document("龙\n你", set())
    assert _lines(doc) == ["龙", "你"]


def test_line_spacing_applies_to_passage_paragraphs_in_both_modes():
    for mode in (MODE_PINYIN, MODE_FONT):
        doc = build_document("龙\n你", set(), mode=mode, line_spacing=1.5)
        assert [p.paragraph_format.line_spacing for p in doc.paragraphs] == [1.5, 1.5]
    assert build_document("龙", set()).paragraphs[0].paragraph_format.line_spacing is None


def test_line_spacing_is_validated():
    for bad in (0.1, 9):
        with pytest.raises(ValueError, match="Line spacing"):
            build_document("龙", set(), line_spacing=bad)


def test_pinyin_first_only_annotates_each_character_once():
    doc = build_document("龙龙你\n龙你", set(), pinyin_first_only=True)
    assert _ruby_pairs(doc) == [("龙", "lóng"), ("你", "nǐ")]
    assert _lines(doc) == ["龙龙你", "龙你"]  # the text itself is intact
    full = build_document("龙龙你\n龙你", set())
    assert len(_ruby_pairs(full)) == 5


def test_answer_key_ignores_first_only_and_name_line_and_is_titled():
    key = build_answer_key("龙龙", pinyin_first_only=True, name_line=True, title="第一课")
    assert len(_ruby_pairs(key)) == 2
    assert [p.text for p in key.paragraphs][0] == "第一课 (Answer key)"
    assert not any(p.text.startswith("Name:") for p in key.paragraphs)
    assert build_answer_key("龙").paragraphs[0].text == "Answer key"


def test_bundle_applies_the_layout_to_both_copies(tmp_path):
    import io
    import zipfile

    z = zipfile.ZipFile(io.BytesIO(worksheet_bundle("龙龙", set(), title="T", name_line=True,
                                                    pinyin_first_only=True)))
    z.extractall(tmp_path)
    student = Document(str(tmp_path / "student_worksheet.docx"))
    teacher = Document(str(tmp_path / "teacher_answer_key.docx"))
    assert student.paragraphs[0].text == "T" and student.paragraphs[1].text.startswith("Name:")
    assert len(_ruby_pairs(student)) == 1 and len(_ruby_pairs(teacher)) == 2


def test_cli_layout_flags(tmp_path):
    passage = tmp_path / "p.txt"
    passage.write_text("龙龙\n", encoding="utf-8")
    out = tmp_path / "w.docx"
    assert main(["YCT=1", str(passage), "-o", str(out), "--title", "课", "--name-line",
                 "--line-spacing", "1.5", "--first-only"]) == 0
    doc = Document(str(out))
    assert doc.paragraphs[0].text == "课" and doc.paragraphs[1].text.startswith("Name:")
    assert doc.paragraphs[2].paragraph_format.line_spacing == 1.5
    assert len(_ruby_pairs(doc)) == 1


def test_layout_settings_travel_in_the_setup_link():
    params = encode_setup({"YCT": 1}, line_spacing=1.5, name_line=True, first_only=True)
    assert params == {"lists": "YCT:1", "spacing": "1.5", "nameline": "1", "first": "1"}
    decoded = decode_setup(params, LEVELS)
    assert (decoded["line_spacing"], decoded["name_line"], decoded["first_only"]) == (1.5, True, True)
    assert "spacing" not in encode_setup({"YCT": 1})  # defaults stay out of the address


def test_decode_clamps_or_ignores_bad_layout_values():
    assert decode_setup({"spacing": "9"}, LEVELS) == {"line_spacing": 2.0}
    assert decode_setup({"spacing": "0.2"}, LEVELS) == {"line_spacing": 1.0}
    for bad in ("abc", "nan", ""):
        assert decode_setup({"spacing": bad}, LEVELS) == {}
    assert decode_setup({"nameline": "yes", "first": "0"}, LEVELS) == {}
